from flask import Blueprint, render_template, request, redirect, url_for, session, flash, send_file
from datetime import datetime
import os
import openpyxl
import logging
from werkzeug.utils import secure_filename
from db import get_db
from utils import log_action, login_required, transliterate_ukrainian, generate_english_name
from gen_docx import gen_doc
import sqlite3

students_bp = Blueprint('students', __name__)

@students_bp.route('/students', methods=['GET'])
@login_required()
def student_list():
    """Отображение списка студентов с поиском, пагинацией, сортировкой и фильтрацией по группе."""
    # Получение параметров запроса
    search = request.args.get('search', '')
    group_id = request.args.get('group_id', type=int)
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 10, type=int)
    sort_by = request.args.get('sort_by', 'id')
    sort_order = request.args.get('sort_order', 'desc')

    # Проверка допустимых значений для пагинации и сортировки
    if per_page not in [10, 20, 50, 100]:
        per_page = 10

    allowed_sort_fields = ['id', 'last_name_UA', 'first_name_UA', 'middle_name_UA', 'birth_date', 'group_id']
    if sort_by not in allowed_sort_fields:
        sort_by = 'id'
    if sort_order not in ['asc', 'desc']:
        sort_order = 'desc'

    offset = (page - 1) * per_page

    # Подключение к базе данных
    conn = get_db()
    role = session.get('role')
    user_id = session.get('user_id')

    # Получение group_ids из базы данных для пользователя
    group_ids = [row['group_id'] for row in conn.execute("SELECT group_id FROM user_groups WHERE user_id = ?", (user_id,)).fetchall()]

    # Логирование действия
    log_action(session.get('username', 'невідомо'), f"переглянув список студентів (group_id={group_id})", group_ids=[group_id] if group_id else group_ids)

    # Базовый запрос для получения студентов (без избыточного WHERE)
    base_query = """
        SELECT s.*, 
               m.id AS has_military,
               m.registration_number_of_the_DRPVR,
               m.military_registration_document,
               m.issued_VOD,
               m.military_accounting_specialty_number,
               m.military_rank,
               m.change_credentials,
               m.reason_for_changing_credentials,
               m.being_on_military_registration,
               m.address_of_residence,
               g.name || ' (' || g.start_year || ', ' || g.study_form || ', ' || g.program_credits || ' кредитів)' AS group_name,
               g.study_form,
               g.program_credits,
               g.qualification_name,
               g.degree_level,
               g.specialty,
               g.educational_program,
               g.knowledge_area,
               g.qualification_name_en,
               g.degree_level_en,
               g.specialty_en,
               g.educational_program_en,
               g.knowledge_area_en
        FROM students s
        LEFT JOIN military m ON m.student_id = s.id
        LEFT JOIN groups g ON s.group_id = g.id
    """
    count_query = "SELECT COUNT(*) FROM students s"
    where_clauses = ["s.archived = FALSE"]  # Базовое условие для исключения архивных студентов
    params = []

    # Фильтр по группе из параметра group_id
    if group_id:
        where_clauses.append("s.group_id = ?")
        params.append(group_id)

    # Ограничение по группе для не-администраторов
    if role != 'admin' and not group_id:
        if group_ids:
            placeholders = ','.join('?' for _ in group_ids)
            where_clauses.append(f"s.group_id IN ({placeholders})")
            params.extend(group_ids)
        else:
            # Если нет групп, возвращаем пустой список
            students = []
            total_students = 0
            students_with_filled_fields = []
            conn.close()
            return render_template(
                'students.html',
                students=students_with_filled_fields,
                search=search,
                group_id=group_id,
                page=page,
                per_page=per_page,
                total_pages=0,
                sort_by=sort_by,
                sort_order=sort_order
            )

    # Проверка доступа к группе для не-администраторов
    if role != 'admin' and group_id and group_id not in group_ids:
        conn.close()
        flash("У вас немає доступу до цієї групи.", "error")
        return redirect(url_for('students.student_list'))

    # Поиск по имени, фамилии, отчеству
    if search:
        where_clauses.append("(s.last_name_UA LIKE ? OR s.first_name_UA LIKE ? OR s.middle_name_UA LIKE ?)")
        params.extend([f'%{search}%', f'%{search}%', f'%{search}%'])

    # Дополнительное условие для корректной сортировки по дате рождения
    if sort_by == 'birth_date':
        where_clauses.append("LENGTH(s.birth_date) = 10 AND INSTR(s.birth_date, '.') = 3 AND INSTR(SUBSTR(s.birth_date, 4), '.') = 3")

    # Формирование условий WHERE
    if where_clauses:
        where_sql = " WHERE " + " AND ".join(where_clauses)
        base_query += where_sql
        count_query += where_sql

    # Сортировка
    if sort_by == 'birth_date':
        base_query += f" ORDER BY SUBSTR(s.birth_date, 7, 4) || SUBSTR(s.birth_date, 4, 2) || SUBSTR(s.birth_date, 1, 2) {sort_order.upper()}"
    elif sort_by in ['last_name_UA', 'first_name_UA', 'middle_name_UA']:
        base_query += f" ORDER BY s.{sort_by} COLLATE UKRAINIAN {sort_order.upper()}"
    else:
        base_query += f" ORDER BY s.{sort_by} {sort_order.upper()}"

    base_query += " LIMIT ? OFFSET ?"
    params_with_limit = params + [per_page, offset]

    # Выполнение запросов
    students = conn.execute(base_query, params_with_limit).fetchall()
    total_students = conn.execute(count_query, params).fetchone()[0]

    # Обработка данных студентов
    students_with_filled_fields = []
    student_fields = ['last_name_UA', 'first_name_UA', 'middle_name_UA', 'last_name_ENG', 'first_name_ENG', 'birth_date', 'group_id', 'edebo_code']
    military_fields = [
        'registration_number_of_the_DRPVR',
        'military_registration_document',
        'issued_VOD',
        'military_accounting_specialty_number',
        'military_rank',
        'change_credentials',
        'reason_for_changing_credentials',
        'being_on_military_registration',
        'address_of_residence'
    ]
    student_total_fields = len(student_fields)
    military_total_fields = len(military_fields)

    for student in students:
        student_dict = dict(student)

        # Вычисление заполненности полей студента
        student_filled_fields = sum(1 for field in student_fields if student_dict.get(field) and str(student_dict.get(field)).strip())
        student_dict['filled_fields'] = student_filled_fields
        student_dict['total_fields'] = student_total_fields

        # Вычисление заполненности военных данных
        if student_dict.get('has_military'):
            military_filled_fields = sum(1 for field in military_fields if student_dict.get(field) and str(student_dict.get(field)).strip())
            student_dict['military_filled_fields'] = military_filled_fields
            student_dict['military_total_fields'] = military_total_fields
        else:
            student_dict['military_filled_fields'] = 0
            student_dict['military_total_fields'] = military_total_fields

        # Вычисление заполненности оценок по предметам
        subjects = conn.execute("SELECT id FROM subjects WHERE group_id = ?", (student_dict['group_id'],)).fetchall()
        student_dict['has_grades'] = len(subjects) > 0
        grades = conn.execute("SELECT grade FROM grades WHERE student_id = ?", (student_dict['id'],)).fetchall()
        grades_filled = sum(1 for grade in grades if grade['grade'] is not None and str(grade['grade']).strip())
        grades_total = len(subjects)
        student_dict['grades_filled'] = grades_filled
        student_dict['grades_total'] = grades_total
        student_dict['grades_fill_percentage'] = (grades_filled / grades_total * 100) if grades_total > 0 else 0

        # Вычисление заполненности активностей (практики, курсовые работы, аттестации)
        practices = conn.execute("SELECT id FROM practices WHERE group_id = ?", (student_dict['group_id'],)).fetchall()
        courseworks = conn.execute("SELECT id FROM courseworks WHERE group_id = ?", (student_dict['group_id'],)).fetchall()
        attestations = conn.execute("SELECT id FROM attestations WHERE group_id = ?", (student_dict['group_id'],)).fetchall()
        student_dict['has_activities'] = len(practices) > 0 or len(courseworks) > 0 or len(attestations) > 0
        activities_grades = conn.execute(
            "SELECT grade FROM activity_grades WHERE student_id = ? AND entity_type IN ('practice', 'coursework', 'attestation')",
            (student_dict['id'],)
        ).fetchall()
        activities_total = len(practices) + len(courseworks) + len(attestations)
        activities_filled = sum(1 for grade in activities_grades if grade['grade'] is not None and str(grade['grade']).strip())
        student_dict['activities_filled_fields'] = activities_filled
        student_dict['activities_total_fields'] = activities_total
        student_dict['activities_fill_percentage'] = (activities_filled / activities_total * 100) if activities_total > 0 else 0

        students_with_filled_fields.append(student_dict)
    groups = conn.execute("""
    SELECT id, name, start_year, study_form 
    FROM groups 
    WHERE archived = FALSE
    ORDER BY start_year DESC, name
    """).fetchall()


    conn.close()

    # Вычисление общего количества страниц
    total_pages = (total_students + per_page - 1) // per_page

    return render_template(
        'students.html',
        students=students_with_filled_fields,
        search=search,
        group_id=group_id,
        groups=groups,
        page=page,
        per_page=per_page,
        total_pages=total_pages,
        sort_by=sort_by,
        sort_order=sort_order
    )
 
@students_bp.route('/students/<int:student_id>')
@login_required('')
def student_details(student_id):
    """Отображение детальной информации о студенте, включая оценки, практики, курсовые, аттестации и документы об образовании."""
    conn = get_db()
    student = conn.execute("""
        SELECT s.*, g.name || ' (' || g.start_year || ', ' || g.study_form || ', ' || g.program_credits || ' кредитів)' AS group_name
        FROM students s
        LEFT JOIN groups g ON s.group_id = g.id
        WHERE s.id = ?
    """, (student_id,)).fetchone()
    
    if not student:
        conn.close()
        flash("Студента не знайдено")
        return redirect(url_for('students.student_list'))
    
    if session.get('role') != 'admin' and student['group_id'] not in session.get('group_ids', []):
        conn.close()
        flash("Доступ заборонено: студент не належить до вашої групи")
        return redirect(url_for('students.student_list'))
    
    military = conn.execute("SELECT * FROM military WHERE student_id = ?", (student_id,)).fetchone()
    
    # Получение оценок и предметов
    grades = conn.execute("""
        SELECT g.grade, s.code, s.name, s.type
        FROM grades g
        JOIN subjects s ON g.subject_id = s.id
        WHERE g.student_id = ?
        ORDER BY s.position
    """, (student_id,)).fetchall()
    
    subjects = conn.execute("""
        SELECT id, code, name, type
        FROM subjects
        WHERE group_id = ?
        ORDER BY position
    """, (student['group_id'],)).fetchall()
    
    grades_dict = {grade['code']: dict(grade) for grade in grades}
    subject_grades = [
        {
            'code': subject['code'],
            'name': subject['name'],
            'type': subject['type'],
            'grade': grades_dict.get(subject['code'], {}).get('grade', None)
        }
        for subject in subjects
    ]
    
    # Получение практик
    practices = conn.execute("""
        SELECT id, code, name, type
        FROM practices
        WHERE group_id = ?
        ORDER BY position
    """, (student['group_id'],)).fetchall()
    
    practice_grades = conn.execute("""
        SELECT ag.grade, p.code, p.name, p.type
        FROM activity_grades ag
        JOIN practices p ON ag.entity_id = p.id
        WHERE ag.student_id = ? AND ag.entity_type = 'practice'
        ORDER BY p.position
    """, (student_id,)).fetchall()
    
    practice_grades_dict = {grade['code']: dict(grade) for grade in practice_grades}
    practice_data = [
        {
            'code': practice['code'],
            'name': practice['name'],
            'type': practice['type'],
            'grade': practice_grades_dict.get(practice['code'], {}).get('grade', None)
        }
        for practice in practices
    ]
    
    # Получение курсовых работ
    courseworks = conn.execute("""
        SELECT id, code, name, type
        FROM courseworks
        WHERE group_id = ?
        ORDER BY position
    """, (student['group_id'],)).fetchall()
    
    coursework_grades = conn.execute("""
        SELECT ag.grade, c.code, c.name, c.type
        FROM activity_grades ag
        JOIN courseworks c ON ag.entity_id = c.id
        WHERE ag.student_id = ? AND ag.entity_type = 'coursework'
        ORDER BY c.position
    """, (student_id,)).fetchall()
    
    coursework_grades_dict = {grade['code']: dict(grade) for grade in coursework_grades}
    coursework_data = [
        {
            'code': coursework['code'],
            'name': coursework['name'],
            'type': coursework['type'],
            'grade': coursework_grades_dict.get(coursework['code'], {}).get('grade', None)
        }
        for coursework in courseworks
    ]
    
    # Получение аттестаций
    attestations = conn.execute("""
        SELECT id, code, name, type
        FROM attestations
        WHERE group_id = ?
        ORDER BY position
    """, (student['group_id'],)).fetchall()
    
    attestation_grades = conn.execute("""
        SELECT ag.grade, a.code, a.name, a.type, ag.name AS student_name
        FROM activity_grades ag
        JOIN attestations a ON ag.entity_id = a.id
        WHERE ag.student_id = ? AND ag.entity_type = 'attestation'
        ORDER BY a.position
    """, (student_id,)).fetchall()
    
    attestation_grades_dict = {grade['code']: dict(grade) for grade in attestation_grades}
    attestation_data = [
        {
            'code': attestation['code'],
            'name': attestation['name'],
            'type': attestation['type'],
            'grade': attestation_grades_dict.get(attestation['code'], {}).get('grade', None),
            'student_name': attestation_grades_dict.get(attestation['code'], {}).get('student_name', None)
        }
        for attestation in attestations
    ]
    
    # Получение документов об образовании
    education_docs = conn.execute("""
        SELECT ed.id, ed.document_type, ed.document_type_en, ed.document_number, ed.institution_name, ed.institution_name_en, 
               ed.country, ed.country_en, ed.completion_date,
               fed.reference_number, fed.reference_institution, fed.reference_institution_en, fed.reference_country, 
               fed.reference_country_en, fed.reference_issue_date, fed.recognition_certificate_number, fed.recognition_issuer, 
               fed.recognition_issuer_en, fed.recognition_date
        FROM education_documents ed
        LEFT JOIN foreign_education_docs fed ON ed.id = fed.education_doc_id
        WHERE ed.student_id = ?
        ORDER BY ed.completion_date DESC
    """, (student_id,)).fetchall()
    
    student_dict = dict(student)
    military_dict = dict(military) if military else None
    
    log_action(session.get('username', 'невідомо'), f"переглянув сторінку студента ID {student_id}", group_ids=[student['group_id']])
    conn.close()
    
    return render_template(
        'student_details.html',
        student=student_dict,
        military=military_dict,
        subject_grades=subject_grades,
        practice_data=practice_data,
        coursework_data=coursework_data,
        attestation_data=attestation_data,
        education_docs=education_docs  # Передаем документы в шаблон
    )

@students_bp.route('/students/add', methods=['GET', 'POST'])
@login_required('')
def add_student():
    """Добавление нового студента."""
    conn = get_db()
    role = session.get('role')
    group_ids = session.get('group_ids', [])
    # Получение доступных групп
    if role == 'admin':
        groups = conn.execute("""
            SELECT id, name, start_year, study_form, program_credits,
                   name || ' (' || start_year || ', ' || study_form || ', ' || program_credits || ' кредитів)' AS display_name
            FROM groups
            WHERE archived = FALSE
            ORDER BY name, start_year
        """).fetchall()
    else:
        placeholders = ','.join('?' for _ in group_ids)
        groups = conn.execute(f"""
            SELECT id, name, start_year, study_form, program_credits,
                   name || ' (' || start_year || ', ' || study_form || ', ' || program_credits || ' кредитів)' AS display_name
            FROM groups
            WHERE id IN ({placeholders}) AND archived = FALSE
            ORDER BY name, start_year
        """, group_ids).fetchall()

    if request.method == 'POST':
        group = request.form.get('group_id')
        try:
            group_int = int(group)
        except (ValueError, TypeError):
            flash("Некоректна група", "error")
            conn.close()
            return render_template('add_student.html', groups=groups)
        if role != 'admin' and group_int not in group_ids:
            flash("Доступ заборонено: група не належить до ваших груп", "error")
            conn.close()
            return render_template('add_student.html', groups=groups)

        # Валидация даты рождения
        birth_date_raw = request.form['birth_date'].strip()
        birth_date_clean = birth_date_raw.replace("-", ".")
        try:
            datetime.strptime(birth_date_clean, "%d.%m.%Y")
            birth_date = birth_date_clean
        except ValueError:
            flash("Невірний формат дати. Введіть у форматі ДД.ММ.РРРР")
            conn.close()
            return render_template('add_student.html', groups=groups)

        # Генерация английских имен (предполагаем, что функция generate_english_name определена)
        last_name_ua = request.form['last_name_UA']
        first_name_ua = request.form['first_name_UA']
        last_name_eng, first_name_eng = generate_english_name(last_name_ua, first_name_ua)

        # Вставка студента
        conn.execute("""
            INSERT INTO students (
                last_name_UA, first_name_UA, middle_name_UA,
                last_name_ENG, first_name_ENG, birth_date, group_id, edebo_code
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            last_name_ua,
            first_name_ua,
            request.form.get('middle_name_UA'),
            last_name_eng,
            first_name_eng,
            birth_date,
            group_int,
            request.form.get('edebo_code')
        ))
        conn.commit()
        student_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]

        # Валидация и обработка военных данных
        issued_VOD_raw = request.form.get('issued_VOD', '').strip()
        if issued_VOD_raw:  # Только если есть данные, выполняем валидацию
            issued_VOD_clean = issued_VOD_raw.replace("-", ".")
            try:
                datetime.strptime(issued_VOD_clean, "%d.%m.%Y")
                issued_VOD = issued_VOD_clean
            except ValueError:
                flash("Невірний формат дати видачі ВОД. Введіть у форматі ДД.ММ.РРРР")
                conn.close()
                return render_template('add_student.html', groups=groups, student_id=student_id)
        else:
            issued_VOD = None  # Если поле пустое, устанавливаем NULL

        # Проверка наличия хотя бы одного поля военных данных
        military_fields = [
            'military_registration_document',
            'registration_number_of_the_DRPVR',
            'military_accounting_specialty_number',
            'military_rank',
            'change_credentials',
            'reason_for_changing_credentials',
            'being_on_military_registration',
            'address_of_residence'
        ]
        if any(request.form.get(field) for field in military_fields):
            conn.execute("""
                INSERT INTO military (
                    student_id,
                    registration_number_of_the_DRPVR,
                    military_registration_document,
                    issued_VOD,
                    military_accounting_specialty_number,
                    military_rank,
                    change_credentials,
                    reason_for_changing_credentials,
                    being_on_military_registration,
                    address_of_residence
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                student_id,
                request.form.get('registration_number_of_the_DRPVR'),
                request.form.get('military_registration_document'),
                issued_VOD,
                request.form.get('military_accounting_specialty_number'),
                request.form.get('military_rank'),
                request.form.get('change_credentials'),
                request.form.get('reason_for_changing_credentials'),
                request.form.get('being_on_military_registration'),
                request.form.get('address_of_residence')
            ))
            conn.commit()

        log_action(
            session.get('username', 'невідомо'),
            f"додав студента: {last_name_ua}",
            group_ids=[group_int]
        )
        conn.close()
        return redirect(url_for('students.student_list'))

    conn.close()
    log_action(session.get('username', 'невідомо'), "відкрив форму додавання студента", group_ids=session.get('group_ids', []))
    return render_template('add_student.html', groups=groups)
    
@students_bp.route('/students/<int:student_id>/edit', methods=['GET', 'POST'])
@login_required('')
def edit_student(student_id):
    """Редактирование данных студента."""
    conn = get_db()
    role = session.get('role')
    group_ids = session.get('group_ids', [])

    student = conn.execute("SELECT * FROM students WHERE id = ?", (student_id,)).fetchone()
    if not student:
        flash('Студента не знайдено', 'error')
        conn.close()
        return redirect(url_for('students.student_list'))

    if role != 'admin' and student['group_id'] not in group_ids:
        flash('Ви не маєте доступу до цього студента', 'error')
        conn.close()
        return redirect(url_for('students.student_list'))

    # Доступные группы
    if role == 'admin':
        groups = conn.execute("""
            SELECT id, name, start_year, study_form, program_credits,
                   name || ' (' || start_year || ', ' || study_form || ', ' || program_credits || ' кредитів)' AS display_name
            FROM groups
            WHERE archived = FALSE
            ORDER BY name, start_year
        """).fetchall()
    else:
        placeholders = ','.join('?' for _ in group_ids)
        groups = conn.execute(f"""
            SELECT id, name, start_year, study_form, program_credits,
                   name || ' (' || start_year || ', ' || study_form || ', ' || program_credits || ' кредитів)' AS display_name
            FROM groups
            WHERE id IN ({placeholders})
            ORDER BY name, start_year
        """, group_ids).fetchall()

    if request.method == 'POST':
        group = request.form.get('group_id')
        try:
            group_int = int(group)
        except (ValueError, TypeError):
            flash("Некоректна група", "error")
            conn.close()
            return render_template('edit_student.html', student=student, groups=groups)

        if role != 'admin' and group_int not in group_ids:
            flash("Доступ заборонено: група не належить до ваших груп", "error")
            conn.close()
            return render_template('edit_student.html', student=student, groups=groups)

        if 'update_english_names' in request.form:
            # Обновление английских имен
            last_name_ua = request.form['last_name_UA']
            first_name_ua = request.form['first_name_UA']
            last_name_eng, first_name_eng = generate_english_name(last_name_ua, first_name_ua)
            conn.execute("""
                UPDATE students 
                SET last_name_ENG = ?, first_name_ENG = ?
                WHERE id = ?
            """, (last_name_eng, first_name_eng, student_id))
            conn.commit()
            flash('Англійські імена оновлено.', 'success')
            conn.close()
            return redirect(url_for('students.edit_student', student_id=student_id))
        else:
            # Обычное сохранение данных
            birth_date_raw = request.form['birth_date'].strip()
            birth_date_clean = birth_date_raw.replace("-", ".")
            try:
                datetime.strptime(birth_date_clean, "%d.%m.%Y")
                birth_date = birth_date_clean
            except ValueError:
                flash("Невірний формат дати. Введіть у форматі ДД.ММ.РРРР")
                conn.close()
                return render_template('edit_student.html', student=student, groups=groups)

            conn.execute("""
                UPDATE students SET
                    last_name_UA=?, first_name_UA=?, middle_name_UA=?,
                    last_name_ENG=?, first_name_ENG=?, birth_date=?,
                    group_id=?, edebo_code=?
                WHERE id=?
            """, (
                request.form['last_name_UA'],
                request.form['first_name_UA'],
                request.form.get('middle_name_UA'),
                request.form.get('last_name_ENG'),
                request.form.get('first_name_ENG'),
                birth_date,
                group_int,
                request.form.get('edebo_code'),
                student_id
            ))
            conn.commit()

        log_action(session.get('username', 'невідомо'), f"редагував студента ID {student_id}", group_ids=[group_int])
        conn.close()
        return redirect(url_for('students.student_list'))

    conn.close()
    return render_template('edit_student.html', student=student, groups=groups)
  
@students_bp.route('/students/<int:student_id>/delete')
@login_required('admin')
def delete_student(student_id):
    """Удаление студента по его ID."""
    conn = get_db()
    student = conn.execute("SELECT last_name_UA, group_id FROM students WHERE id = ?", (student_id,)).fetchone()
    if student:
        conn.execute("DELETE FROM military WHERE student_id = ?", (student_id,))
        conn.execute("DELETE FROM grades WHERE student_id = ?", (student_id,))
        conn.execute("DELETE FROM activity_grades WHERE student_id = ?", (student_id,))
        conn.execute("DELETE FROM students WHERE id = ?", (student_id,))
        conn.commit()
        log_action(session.get('username', 'невідомо'), f"видалив студента ID {student_id}: {student['last_name_UA']}", group_ids=[student['group_id']])
    else:
        flash("Студента не знайдено")
    conn.close()
    return redirect(url_for('students.student_list'))

@students_bp.route('/students/<int:student_id>/military/add', methods=['GET', 'POST'])
@login_required('')
def add_military(student_id):
    """Добавление военных данных для студента."""
    if request.method == 'POST':
        issued_VOD_raw = request.form.get('issued_VOD', '').strip()  # Получаем значение с дефолтом пустой строки

        # Проверяем, пустое ли поле
        if issued_VOD_raw:  # Только если есть данные, выполняем валидацию
            issued_VOD_clean = issued_VOD_raw.replace("-", ".")
            try:
                datetime.strptime(issued_VOD_clean, "%d.%m.%Y")
                issued_VOD = issued_VOD_clean
            except ValueError:
                flash("Невірний формат дати. Введіть у форматі ДД.ММ.РРРР")
                return render_template('add_military.html', student_id=student_id)
        else:
            issued_VOD = None  # Если поле пустое, устанавливаем NULL

        data = (
            student_id,
            request.form['registration_number_of_the_DRPVR'],
            request.form['military_registration_document'],
            issued_VOD,
            request.form['military_accounting_specialty_number'],
            request.form['military_rank'],
            request.form['change_credentials'],
            request.form['reason_for_changing_credentials'],
            request.form['being_on_military_registration'],
            request.form['address_of_residence'],
        )

        conn = get_db()
        conn.execute("""
            INSERT INTO military (
                student_id, registration_number_of_the_DRPVR, military_registration_document,
                issued_VOD, military_accounting_specialty_number, military_rank,
                change_credentials, reason_for_changing_credentials,
                being_on_military_registration, address_of_residence
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, data)
        conn.commit()
        conn.close()

        log_action(session.get('username', 'невідомо'), f"додав військові дані для студента ID {student_id}", session.get('group_id'))
        return redirect(url_for('students.student_list'))

    log_action(session.get('username', 'невідомо'), f"відкрив форму додавання військових даних для студента ID {student_id}", session.get('group_id'))
    return render_template('add_military.html', student_id=student_id)

@students_bp.route('/students/<int:student_id>/military', methods=['GET', 'POST'])
@login_required('')
def military_data(student_id):
    """Редактирование или добавление военных данных студента."""
    conn = get_db()
    military = conn.execute("SELECT * FROM military WHERE student_id = ?", (student_id,)).fetchone()

    if request.method == 'POST':
        issued_VOD_raw = request.form['issued_VOD'].strip()
        issued_VOD_clean = issued_VOD_raw.replace("-", ".")

        try:
            datetime.strptime(issued_VOD_clean, "%d.%m.%Y")
            issued_VOD = issued_VOD_clean
        except ValueError:
            flash("Невірний формат дати. Введіть у форматі ДД.ММ.РРРР")
            return render_template('edit_military.html', student_id=student_id, military=military)
        
        data = (
            request.form['registration_number_of_the_DRPVR'],
            request.form['military_registration_document'],
            issued_VOD,
            request.form['military_accounting_specialty_number'],
            request.form['military_rank'],
            request.form['change_credentials'],
            request.form['reason_for_changing_credentials'],
            request.form['being_on_military_registration'],
            request.form['address_of_residence'],
            student_id
        )
        if military:
            conn.execute("""
                UPDATE military SET
                    registration_number_of_the_DRPVR=?,
                    military_registration_document=?,
                    issued_VOD=?,
                    military_accounting_specialty_number=?,
                    military_rank=?,
                    change_credentials=?,
                    reason_for_changing_credentials=?,
                    being_on_military_registration=?,
                    address_of_residence=?
                WHERE student_id=?
            """, data)
        else:
            conn.execute("""
                INSERT INTO military (
                    registration_number_of_the_DRPVR, military_registration_document,
                    issued_VOD, military_accounting_specialty_number, military_rank,
                    change_credentials, reason_for_changing_credentials,
                    being_on_military_registration, address_of_residence, student_id
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, data)
        conn.commit()
        log_action(session.get('username', 'невідомо'), f"змінив військові дані для студента ID {student_id}", session.get('group_id'))
        conn.close()
        return redirect(url_for('students.student_list'))

    conn.close()
    log_action(session.get('username', 'невідомо'), f"відкрив форму редагування військових даних для студента ID {student_id}", session.get('group_id'))
    return render_template('edit_military.html', student_id=student_id, military=military)

@students_bp.route('/students/<int:student_id>/military/delete')
@login_required('admin')
def delete_military(student_id):
    """Удаление военных данных студента."""
    conn = get_db()
    conn.execute("DELETE FROM military WHERE student_id = ?", (student_id,))
    conn.commit()
    log_action(session.get('username', 'невідомо'), f"видалив військові дані для студента ID {student_id}")
    conn.close()
    return redirect(url_for('students.student_list'))

@students_bp.route('/students/<int:student_id>/generate', methods=['GET', 'POST'])
@login_required()
def generate(student_id):
    """Генерация документа для студента."""
    if request.method == 'POST':
        selected_template = request.form.get('template', 'template.docx')
        conn = get_db()
        conn.row_factory = sqlite3.Row
        student = conn.execute("""
            SELECT s.*, 
                   g.name || ' (' || g.start_year || ', ' || g.study_form || ', ' || g.program_credits || ' кредитів)' AS group_name,
                   g.start_year,
                   g.study_form,
                   g.program_credits,
                   g.qualification_name,
                   g.degree_level,
                   g.specialty,
                   g.educational_program,
                   g.knowledge_area,
                   g.qualification_name_en,
                   g.degree_level_en,
                   g.entry_requirements,
                   g.entry_requirements_en,
                   g.learning_outcomes,
                   g.learning_outcomes_en,
                   g.program_includes,
                   g.program_includes_en,
                   g.specialty_en,
                   g.educational_program_en,
                   g.knowledge_area_en,
                   g.institution_name_and_status,
                   g.institution_name_and_status_en,
                   m.registration_number_of_the_DRPVR,
                   m.military_registration_document,
                   m.issued_VOD,
                   m.military_accounting_specialty_number,
                   m.military_rank,
                   m.change_credentials,
                   m.reason_for_changing_credentials,
                   m.being_on_military_registration,
                   m.address_of_residence
            FROM students s
            LEFT JOIN groups g ON s.group_id = g.id
            LEFT JOIN military m ON s.id = m.student_id
            WHERE s.id = ?
        """, (student_id,)).fetchone()
        military = conn.execute("SELECT * FROM military WHERE student_id=?", (student_id,)).fetchone()
        conn.close()

        if student:
            student_dict = dict(student)
            required_fields = ['last_name_UA', 'first_name_UA', 'id', 'group_id']
            if not all(key in student_dict for key in required_fields):
                logging.error(f"Неповні дані студента ID {student_id}: {student_dict}")
                flash("Дані студента неповні (відсутні необхідні поля)")
                return redirect(url_for('students.student_list'))

            military_dict = dict(military) if military else {}
            student_name_part = f"{student_dict['last_name_UA']}_{student_dict['first_name_UA']}".replace(" ", "_")
            project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            output_dir = os.path.join(project_root, 'generated_docs')
            os.makedirs(output_dir, exist_ok=True)
            output_path = os.path.join(output_dir, f"{student_name_part}.docx")

            try:
                gen_doc(student_dict, military_dict, template=selected_template, out=output_path, user_name=session.get('username', 'невідомо'))
            except Exception as e:
                logging.error(f"Помилка при генерації документа для студента ID {student_id}: {str(e)}")
                flash(f"Помилка при генерації документа: {str(e)}")
                return redirect(url_for('students.student_list'))

            log_action(session.get('username', 'невідомо'), f"згенерував документ для студента ID {student_id}", session.get('group_id'))
            try:
                return send_file(output_path, as_attachment=True)
            except Exception as e:
                logging.error(f"Помилка при відправленні файлу {output_path}: {str(e)}")
                flash(f"Помилка при відправленні файлу: {str(e)}")
                return redirect(url_for('students.student_list'))

        flash("Студента не знайдено")
        return redirect(url_for('students.student_list'))

    log_action(session.get('username', 'невідомо'), f"відкрив форму генерації документа для студента ID {student_id}", session.get('group_id'))
    return render_template('generate_word.html', student_id=student_id)

@students_bp.route('/activities_grades/<int:student_id>', methods=['GET', 'POST'])
@login_required()
def edit_activities_grades(student_id):
    """Редактирование оценок студента по практикам, курсовым работам и аттестациям."""
    conn = get_db()

    # Получение данных о студенте
    student = conn.execute("""
        SELECT s.*, 
               g.name || ' (' || g.start_year || ', ' || g.study_form || ', ' || g.program_credits || ' кредитів)' AS group_name,
               g.study_form,
               g.program_credits,
               g.qualification_name,
               g.degree_level,
               g.specialty,
               g.educational_program,
               g.knowledge_area,
               g.qualification_name_en,
               g.degree_level_en,
               g.specialty_en,
               g.educational_program_en,
               g.knowledge_area_en
        FROM students s
        LEFT JOIN groups g ON s.group_id = g.id
        WHERE s.id = ?
    """, (student_id,)).fetchone()

    if not student:
        conn.close()
        flash("Студента не знайдено", "error")
        return redirect(url_for('students.student_list'))

    # Проверка доступа
    if session.get('role') != 'admin' and student['group_id'] not in session.get('group_ids', []):
        conn.close()
        flash("Доступ заборонено: студент не належить до вашої групи", "error")
        return redirect(url_for('students.student_list'))

    # Получение данных о практиках
    practices = conn.execute("""
        SELECT id, code, name, credits, type, position
        FROM practices
        WHERE group_id = ?
        ORDER BY position
    """, (student['group_id'],)).fetchall()

    # Получение данных о курсовых работах
    courseworks = conn.execute("""
        SELECT id, code, name, credits, type, position
        FROM courseworks
        WHERE group_id = ?
        ORDER BY position
    """, (student['group_id'],)).fetchall()

    # Получение данных об аттестациях с индивидуальным названием
    attestations = conn.execute("""
        SELECT a.id, a.code, a.name, a.credits, a.type, a.position, ag.name AS student_name
        FROM attestations a
        LEFT JOIN activity_grades ag ON ag.entity_id = a.id AND ag.entity_type = 'attestation' AND ag.student_id = ?
        WHERE a.group_id = ?
        ORDER BY position
    """, (student_id, student['group_id'])).fetchall()

    # Получение существующих оценок и названий
    existing_grades = conn.execute("""
        SELECT id, entity_id, entity_type, grade, name
        FROM activity_grades
        WHERE student_id = ?
    """, (student_id,)).fetchall()
    grade_map = {(g['entity_id'], g['entity_type']): {'id': g['id'], 'grade': g['grade'], 'name': g['name']} for g in existing_grades}

    if request.method == 'POST':
        try:
            for entity_type, entities in [('practice', practices), ('coursework', courseworks), ('attestation', attestations)]:
                for entity in entities:
                    grade_key = f'grade_{entity_type}_{entity["id"]}'
                    name_key = f'name_{entity_type}_{entity["id"]}' if entity_type == 'attestation' else None
                    grade_value = request.form.get(grade_key)
                    student_name = request.form.get(name_key) if entity_type == 'attestation' else ''

                    if grade_value:
                        try:
                            grade_value = int(grade_value)
                            if not 0 <= grade_value <= 100:
                                flash(f"Некоректна оцінка для {entity['name']}: має бути від 0 до 100", "error")
                                continue

                            key = (entity['id'], entity_type)
                            if key in grade_map:
                                # Обновление существующей записи
                                conn.execute("""
                                    UPDATE activity_grades
                                    SET grade = ?, name = ?
                                    WHERE id = ? AND student_id = ? AND entity_id = ? AND entity_type = ?
                                """, (grade_value, student_name, grade_map[key]['id'], student_id, entity['id'], entity_type))
                            else:
                                # Вставка новой записи
                                conn.execute("""
                                    INSERT INTO activity_grades (student_id, entity_id, entity_type, grade, name)
                                    VALUES (?, ?, ?, ?, ?)
                                """, (student_id, entity['id'], entity_type, grade_value, student_name))
                        except ValueError:
                            # Удаление записи при некорректной оценке
                            conn.execute("""
                                DELETE FROM activity_grades
                                WHERE student_id = ? AND entity_id = ? AND entity_type = ?
                            """, (student_id, entity['id'], entity_type))
                            flash(f"Некоректна оцінка для {entity['name']}: має бути числом", "error")
                    else:
                        # Удаление записи при пустой оценке
                        if key in grade_map:
                            conn.execute("""
                                DELETE FROM activity_grades
                                WHERE id = ? AND student_id = ? AND entity_id = ? AND entity_type = ?
                            """, (grade_map[key]['id'], student_id, entity['id'], entity_type))
                        else:
                            conn.execute("""
                                DELETE FROM activity_grades
                                WHERE student_id = ? AND entity_id = ? AND entity_type = ?
                            """, (student_id, entity['id'], entity_type))
            conn.commit()
            flash("Оцінки успішно збережено", "success")
            log_action(session.get('username', 'невідомо'), f"відредагував оцінки для студента ID {student_id}", [student['group_id']])
            conn.close()
            return redirect(url_for('students.student_list'))
        except Exception as e:
            conn.rollback()
            flash(f"Помилка при збереженні оцінок: {str(e)}", "error")

    conn.close()
    return render_template(
        "edit_activities_grades.html",
        student=student,
        practices=practices,
        courseworks=courseworks,
        attestations=attestations,
        grade_map=grade_map
    )
    
@students_bp.route('/grades/<int:student_id>', methods=['GET', 'POST'])
@login_required()
def edit_grades(student_id):
    """Редактирование оценок студента."""
    conn = get_db()

    student = conn.execute("""
        SELECT s.*, g.name || ' (' || g.start_year || ', ' || g.study_form || ', ' || g.program_credits || ' кредитів)' AS group_name
        FROM students s
        LEFT JOIN groups g ON s.group_id = g.id
        WHERE s.id = ?
    """, (student_id,)).fetchone()

    if not student:
        flash("Студент не знайдений")
        return redirect(url_for('students.student_list'))

    subjects = conn.execute("""
        SELECT * FROM subjects
        WHERE group_id = ?
    """, (student['group_id'],)).fetchall()

    existing_grades = conn.execute("""
        SELECT subject_id, grade FROM grades
        WHERE student_id = ?
    """, (student_id,)).fetchall()
    grade_map = {g['subject_id']: g['grade'] for g in existing_grades}

    if request.method == 'POST':
        for subject in subjects:
            grade_value = request.form.get(f'grade_{subject["id"]}')
            if grade_value:
                if subject["id"] in grade_map:
                    conn.execute("""
                        UPDATE grades SET grade = ?
                        WHERE student_id = ? AND subject_id = ?
                    """, (grade_value, student_id, subject["id"]))
                else:
                    conn.execute("""
                        INSERT INTO grades (student_id, subject_id, grade)
                        VALUES (?, ?, ?)
                    """, (student_id, subject["id"], grade_value))
        conn.commit()
        flash("Оцінки збережено")
        return redirect(url_for('students.student_list'))

    conn.close()
    return render_template("edit_grades.html", student=student, subjects=subjects, grade_map=grade_map)

@students_bp.route('/import', methods=['GET', 'POST'])
@login_required('admin')
def import_from_excel():
    """Импорт студентов из Excel-файла."""
    if request.method == 'POST':
        file = request.files.get('excel_file')
        if not file or not file.filename.endswith('.xlsx'):
            flash("Будь ласка, виберіть файл формату .xlsx")
            return render_template('import_excel.html')

        filename = secure_filename(file.filename)
        filepath = os.path.join('uploads', filename)
        os.makedirs('uploads', exist_ok=True)
        file.save(filepath)

        wb = openpyxl.load_workbook(filepath)
        sheet = wb.active

        conn = get_db()
        inserted = 0
        skipped = 0

        for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            try:
                group_id = row[0]
                full_name = row[1]
                birth_date_raw = row[2]
                edebo_code = row[3] if row[3] else ''  # Добавляем edebo_code из 4-го столбца
                military_data = row[4:13]  # Сдвигаем военные данные на один столбец вправо

                if not full_name:
                    continue
                name_parts = full_name.strip().split()
                if len(name_parts) != 3:
                    flash(f"❗ Невірний формат ПІБ у рядку {i}: '{full_name}'")
                    skipped += 1
                    continue
                last_name, first_name, middle_name = name_parts

                if isinstance(birth_date_raw, datetime):
                    birth_date = birth_date_raw.strftime("%d.%m.%Y")
                else:
                    birth_date = str(birth_date_raw).strip()

                existing = conn.execute("""
                    SELECT id FROM students
                    WHERE last_name_UA=? AND first_name_UA=? AND middle_name_UA=? AND birth_date=?
                """, (last_name, first_name, middle_name, birth_date)).fetchone()
                if existing:
                    skipped += 1
                    continue

                # Генерация английских имен
                last_name_eng, first_name_eng = generate_english_name(last_name, first_name)

                conn.execute("""
                    INSERT INTO students (
                        last_name_UA, first_name_UA, middle_name_UA,
                        last_name_ENG, first_name_ENG, birth_date, group_id, edebo_code
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    last_name, first_name, middle_name,
                    last_name_eng, first_name_eng, birth_date, group_id, edebo_code
                ))
                student_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]

                issued_VOD_raw = military_data[2]
                if isinstance(issued_VOD_raw, datetime):
                    issued_VOD = issued_VOD_raw.strftime("%d.%m.%Y")
                elif isinstance(issued_VOD_raw, str):
                    issued_VOD = issued_VOD_raw.strip().replace('-', '.')
                    try:
                        parsed = datetime.strptime(issued_VOD, "%d.%m.%Y")
                        issued_VOD = parsed.strftime("%d.%m.%Y")
                    except ValueError:
                        issued_VOD = ''
                else:
                    issued_VOD = ''

                conn.execute("""
                    INSERT INTO military (
                        student_id,
                        registration_number_of_the_DRPVR,
                        military_registration_document,
                        issued_VOD,
                        military_accounting_specialty_number,
                        military_rank,
                        change_credentials,
                        reason_for_changing_credentials,
                        being_on_military_registration,
                        address_of_residence
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    student_id,
                    military_data[0],
                    military_data[1],
                    issued_VOD,
                    military_data[3],
                    military_data[4],
                    military_data[5],
                    military_data[6],
                    military_data[7],
                    military_data[8]
                ))

                inserted += 1

            except Exception as e:
                flash(f"⚠️ Помилка в рядку {i}: {e}")
                skipped += 1
                continue

        conn.commit()
        conn.close()

        log_action(session.get('username', 'невідомо'), f"виконав імпорт студентів: додано {inserted}, пропущено {skipped}")
        flash(f"✅ Імпорт завершено. Додано: {inserted}, пропущено: {skipped}")
        return redirect(url_for('students.student_list'))

    log_action(session.get('username', 'невідомо'), "відкрив форму імпорту студентів")
    return render_template('import_excel.html')